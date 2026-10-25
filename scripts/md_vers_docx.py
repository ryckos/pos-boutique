"""Convertit un Markdown simple (titres, listes, tableaux, blocs de code) en document Word.

Sert à tenir la version Word d'un document de docs/ à jour après chaque modification du .md :

    pip install python-docx
    python scripts/md_vers_docx.py docs/DEPLOIEMENT.md docs/DEPLOIEMENT.docx "Procédure de déploiement — Ma Boutique POS"
"""
import re
import sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BLEU = RGBColor(0x1F, 0x5F, 0xBF)
MARINE = RGBColor(0x13, 0x29, 0x4B)


def ombrer(element_pr, couleur):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), couleur)
    element_pr.append(shd)


def texte_riche(paragraphe, texte, taille=None, gras=False):
    """Ajoute le texte en gérant **gras** et `code`."""
    texte = texte.replace('- [ ] ', '☐ ').replace('[ ]', '☐')
    for morceau in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', texte):
        if not morceau:
            continue
        if morceau.startswith('**') and morceau.endswith('**'):
            # Le gras peut contenir du code (**`main`**) : on le traite comme un texte en gras.
            texte_riche(paragraphe, morceau[2:-2], taille, gras=True)
            continue
        elif morceau.startswith('`') and morceau.endswith('`'):
            r = paragraphe.add_run(morceau[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(9)
            r.font.color.rgb = MARINE
            r.bold = gras or None
        else:
            r = paragraphe.add_run(morceau)
            if gras:
                r.bold = True
        if taille:
            r.font.size = Pt(taille)


def nouvelle_numerotation(doc, style):
    """Crée un compteur neuf (départ à 1) pour une liste numérotée, basé sur celui du style."""
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    num_id_style = doc.styles[style].element.pPr.numPr.numId.val
    abstract = numbering.num_having_numId(num_id_style).abstractNumId.val
    num = numbering.add_num(abstract)
    override = num.add_lvlOverride(ilvl=0)
    override.add_startOverride(1)
    return num.numId


def appliquer_numerotation(paragraphe, num_id):
    num_pr = paragraphe._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = 0
    num_pr.get_or_add_numId().val = num_id


def cellules(ligne):
    return [c.strip() for c in ligne.strip().strip('|').split('|')]


def structurelle(ligne):
    """Ligne qui commence un nouveau bloc (titre, liste, tableau, code, séparateur) ou ligne vide."""
    l = ligne.strip()
    return not l or l.startswith(('#', '|', '```', '---')) or re.match(r'^(- |\d+\. )', l) is not None


def recoller(lignes):
    """Recolle les lignes de suite d'un paragraphe ou d'un élément de liste : le gras et le code qui
    courent sur deux lignes du Markdown restent d'un seul tenant."""
    sortie = []
    dans_code = False
    for ligne in lignes:
        if ligne.strip().startswith('```'):
            dans_code = not dans_code
            sortie.append(ligne)
            continue
        precedente = sortie[-1].strip() if sortie else ''
        if (
            not dans_code
            and precedente
            and not structurelle(ligne)
            and not precedente.startswith(('#', '|', '```', '---'))
        ):
            sortie[-1] = sortie[-1].rstrip() + ' ' + ligne.strip()
        else:
            sortie.append(ligne)
    return sortie


def convertir(source, cible, titre_document):
    lignes = recoller(open(source, encoding='utf-8').read().split('\n'))
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    for cote in ('left_margin', 'right_margin'):
        setattr(sec, cote, Cm(2))
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)

    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    for nom, taille in (('Title', 22), ('Heading 1', 15), ('Heading 2', 12.5), ('Heading 3', 11)):
        st = doc.styles[nom]
        st.font.name = 'Calibri'
        st.font.size = Pt(taille)
        st.font.color.rgb = MARINE if nom == 'Title' else BLEU
        st.font.bold = True

    # Pied de page : titre et numéro de page.
    pied = sec.footer.paragraphs[0]
    pied.text = f'{titre_document} — page '
    for r in pied.runs:
        r.font.size = Pt(8)
    run = pied.add_run()
    for balise, texte in (('begin', None), (None, 'PAGE'), ('end', None)):
        if balise:
            fld = OxmlElement('w:fldChar')
            fld.set(qn('w:fldCharType'), balise)
            run._r.append(fld)
        else:
            instr = OxmlElement('w:instrText')
            instr.text = texte
            run._r.append(instr)
    run.font.size = Pt(8)

    i = 0
    paragraphe_ouvert = None
    dernier_item = None
    num_courant = None  # compteur de la liste numérotée en cours
    while i < len(lignes):
        ligne = lignes[i]
        brut = ligne.rstrip()

        # Bloc de code
        if brut.lstrip().startswith("```"):
            dans_liste = brut.startswith("  ") and num_courant is not None
            i += 1
            code = []
            while i < len(lignes) and not lignes[i].lstrip().startswith('```'):
                code.append(lignes[i].rstrip())
                i += 1
            for c in code:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(1.2 if dans_liste else 0.5)
                ombrer(p._p.get_or_add_pPr(), 'EEF2F8')
                r = p.add_run(c or ' ')
                r.font.name = 'Consolas'
                r.font.size = Pt(8.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            i += 1
            paragraphe_ouvert = None
            if not dans_liste:
                dernier_item = num_courant = None
            continue

        # Tableau
        if brut.startswith('|') and i + 1 < len(lignes) and re.match(r'^\|[\s:|-]+\|$', lignes[i + 1].strip()):
            entete = cellules(brut)
            corps = []
            i += 2
            while i < len(lignes) and lignes[i].strip().startswith('|'):
                corps.append(cellules(lignes[i]))
                i += 1
            t = doc.add_table(rows=1 + len(corps), cols=len(entete))
            t.style = 'Table Grid'
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            taille = 8.5 if len(entete) >= 5 else 9.5
            for j, h in enumerate(entete):
                c = t.rows[0].cells[j]
                c.paragraphs[0].text = ''
                texte_riche(c.paragraphs[0], h, taille, gras=True)
                ombrer(c._tc.get_or_add_tcPr(), 'E5EEFB')
            for k, ligne_t in enumerate(corps):
                for j in range(len(entete)):
                    c = t.rows[k + 1].cells[j]
                    texte_riche(c.paragraphs[0], ligne_t[j] if j < len(ligne_t) else '', taille)
            # En-tête répété sur chaque page.
            tr_pr = t.rows[0]._tr.get_or_add_trPr()
            rep = OxmlElement('w:tblHeader')
            rep.set(qn('w:val'), 'true')
            tr_pr.append(rep)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            paragraphe_ouvert = dernier_item = num_courant = None
            continue

        if not brut.strip():
            paragraphe_ouvert = None
            i += 1
            continue

        if brut.strip() == '---':
            paragraphe_ouvert = dernier_item = num_courant = None
            i += 1
            continue

        m = re.match(r'^(#{1,4})\s+(.*)$', brut)
        if m:
            niveau = len(m.group(1))
            texte = m.group(2)
            if niveau == 1:
                p = doc.add_paragraph(style='Title')
                texte_riche(p, texte)
            else:
                # Les annexes commencent sur une nouvelle page.
                if niveau == 2 and texte.startswith('Annexe'):
                    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                p = doc.add_heading(level=niveau - 1)
                texte_riche(p, texte)
            paragraphe_ouvert = dernier_item = num_courant = None
            i += 1
            continue

        m = re.match(r'^(\s*)(- |\d+\. )(.*)$', brut)
        if m:
            retrait = len(m.group(1))
            numerote = m.group(2).strip()[0].isdigit()
            style = 'List Number' if numerote else 'List Bullet'
            if retrait >= 2:
                style += ' 2'
            elif not numerote:
                num_courant = None
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(2)
            if numerote and retrait < 2:
                if num_courant is None:
                    num_courant = nouvelle_numerotation(doc, 'List Number')
                appliquer_numerotation(p, num_courant)
            texte_riche(p, m.group(3))
            dernier_item = p
            paragraphe_ouvert = None
            i += 1
            continue

        # Suite d'un élément de liste (ligne indentée)
        if brut.startswith('  ') and dernier_item is not None:
            texte_riche(dernier_item, ' ' + brut.strip())
            i += 1
            continue

        num_courant = None
        if paragraphe_ouvert is not None:
            texte_riche(paragraphe_ouvert, ' ' + brut.strip())
        else:
            paragraphe_ouvert = doc.add_paragraph()
            texte_riche(paragraphe_ouvert, brut.strip())
        dernier_item = None
        i += 1

    doc.core_properties.title = titre_document
    doc.core_properties.language = 'fr-FR'
    doc.save(cible)


if __name__ == '__main__':
    convertir(sys.argv[1], sys.argv[2], sys.argv[3])
